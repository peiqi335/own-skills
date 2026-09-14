#!/usr/bin/env python3
"""Inspect an editable PPTX package and report structural QA findings."""

from __future__ import annotations

import argparse
from collections import Counter
from io import BytesIO
import json
import math
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import string
import tempfile
from typing import Any, Iterable
import unicodedata
from urllib.parse import unquote
import xml.etree.ElementTree as ET
import zipfile


EMU_PER_INCH = 914400
BLEED_LIMIT_PERCENT = 1.0
MAX_ZIP_ENTRIES = 10_000
MAX_ENTRY_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_TOTAL_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MIN_RATIO_CHECK_BYTES = 1024 * 1024

SLIDE_PART = re.compile(r"ppt/slides/slide\d+\.xml$")
NOTES_PART = re.compile(r"ppt/notesSlides/notesSlide\d+\.xml$")
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
RASTER_EXTENSIONS = {
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
SLIDE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
)
NOTES_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml"
)

Matrix = tuple[float, float, float, float, float, float]
IDENTITY: Matrix = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


def _normalize_controls(value: Any) -> str:
    text = str(value)
    text = "".join(
        " " if unicodedata.category(character).startswith("C") else character
        for character in text
    )
    return re.sub(r"\s+", " ", text).strip()


def _markdown_escape(value: Any) -> str:
    text = _normalize_controls(value)
    specials = set(string.punctuation)
    return "".join(f"\\{character}" if character in specials else character for character in text)


def _empty_result(path: Path) -> dict[str, Any]:
    try:
        display_path = str(path.resolve(strict=False))
    except (OSError, RuntimeError):
        display_path = str(path)
    return {
        "ok": False,
        "file": display_path,
        "slides": 0,
        "media": 0,
        "notes_slides": 0,
        "out_of_bounds": [],
        "warnings": [],
        "errors": [],
    }


def _archive_limit_errors(entries: Iterable[Any]) -> list[str]:
    infos = list(entries)
    errors: list[str] = []
    if len(infos) > MAX_ZIP_ENTRIES:
        errors.append(
            f"ZIP entry count exceeds the safety limit ({len(infos)} > "
            f"{MAX_ZIP_ENTRIES})."
        )

    total = sum(max(0, int(info.file_size)) for info in infos)
    if total > MAX_TOTAL_UNCOMPRESSED_BYTES:
        errors.append(
            "ZIP total uncompressed size exceeds the safety limit "
            f"({total} > {MAX_TOTAL_UNCOMPRESSED_BYTES} bytes)."
        )

    for info in infos:
        size = max(0, int(info.file_size))
        compressed = max(0, int(info.compress_size))
        name = _normalize_controls(info.filename)
        if size > MAX_ENTRY_UNCOMPRESSED_BYTES:
            errors.append(
                f"ZIP entry '{name}' exceeds the uncompressed-size safety limit "
                f"({size} > {MAX_ENTRY_UNCOMPRESSED_BYTES} bytes)."
            )
        if size >= MIN_RATIO_CHECK_BYTES:
            ratio = math.inf if compressed == 0 else size / compressed
            if ratio > MAX_COMPRESSION_RATIO:
                errors.append(
                    f"ZIP entry '{name}' has a suspicious compression ratio "
                    f"({ratio:.1f} > {MAX_COMPRESSION_RATIO})."
                )
    return errors


def _content_types(
    package: zipfile.ZipFile,
) -> tuple[dict[str, str], dict[str, str]]:
    root = ET.fromstring(package.read("[Content_Types].xml"))
    if root.tag != f"{{{CONTENT_TYPES_NS}}}Types":
        raise ValueError("[Content_Types].xml has an unexpected root element")
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for child in root:
        if child.tag == f"{{{CONTENT_TYPES_NS}}}Default":
            extension = child.get("Extension")
            content_type = child.get("ContentType")
            if extension and content_type:
                defaults[extension.lower()] = content_type
        elif child.tag == f"{{{CONTENT_TYPES_NS}}}Override":
            part_name = child.get("PartName")
            content_type = child.get("ContentType")
            if part_name and content_type:
                overrides[part_name.lstrip("/")] = content_type
    return defaults, overrides


def _content_type_for(
    part: str, defaults: dict[str, str], overrides: dict[str, str]
) -> str | None:
    if part in overrides:
        return overrides[part]
    suffix = PurePosixPath(part).suffix.lstrip(".").lower()
    return defaults.get(suffix)


def _validate_content_type_coverage(
    names: set[str],
    defaults: dict[str, str],
    overrides: dict[str, str],
) -> tuple[dict[str, str], list[str]]:
    resolved: dict[str, str] = {}
    errors: list[str] = []
    for part in sorted(names):
        expected: str | None = None
        if SLIDE_PART.fullmatch(part):
            expected = SLIDE_CONTENT_TYPE
        elif NOTES_PART.fullmatch(part):
            expected = NOTES_CONTENT_TYPE
        elif part == "ppt/presentation.xml":
            expected = (
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation.main+xml"
            )
        elif part.startswith("ppt/media/"):
            expected = "media"
        else:
            continue

        actual = _content_type_for(part, defaults, overrides)
        if actual is not None:
            resolved[part] = actual
        if actual is None:
            errors.append(f"No content type covers package part '{part}'.")
        elif expected == "media" and actual in {"application/xml", "text/xml"}:
            errors.append(f"Media part '{part}' has an invalid content type '{actual}'.")
        elif expected not in {None, "media"} and actual != expected:
            errors.append(
                f"Package part '{part}' has content type '{actual}', expected "
                f"'{expected}'."
            )
    return resolved, errors


def _relationship_source(rels_path: str) -> tuple[str | None, str]:
    if rels_path == "_rels/.rels":
        return None, ""
    path = PurePosixPath(rels_path)
    if path.parent.name != "_rels" or not path.name.endswith(".rels"):
        raise ValueError(f"Malformed relationships part path '{rels_path}'.")
    source_dir = path.parent.parent
    source = (source_dir / path.name[:-5]).as_posix()
    return source, source_dir.as_posix()


def _resolve_internal_target(source_dir: str, target: str) -> str:
    decoded = unquote(target.split("#", 1)[0])
    if not decoded or "?" in decoded:
        raise ValueError(f"Invalid internal relationship target '{target}'.")
    if decoded.startswith("/"):
        normalized = posixpath.normpath(decoded.lstrip("/"))
    else:
        normalized = posixpath.normpath(posixpath.join(source_dir, decoded))
    if normalized in {"", ".", ".."} or normalized.startswith("../"):
        raise ValueError(f"Internal relationship target escapes the package: '{target}'.")
    return normalized


def _relationship_target_compatible(
    kind: str, target: str, content_type: str | None
) -> tuple[bool, str | None]:
    embedded_media = target.startswith("ppt/media/")
    if kind == "slide":
        compatible = (
            bool(SLIDE_PART.fullmatch(target))
            and content_type == SLIDE_CONTENT_TYPE
        )
        expectation = f"a slide part with content type '{SLIDE_CONTENT_TYPE}'"
    elif kind == "notesSlide":
        compatible = (
            bool(NOTES_PART.fullmatch(target))
            and content_type == NOTES_CONTENT_TYPE
        )
        expectation = f"a notesSlide part with content type '{NOTES_CONTENT_TYPE}'"
    elif kind == "image":
        compatible = embedded_media and bool(
            content_type and content_type.startswith("image/")
        )
        expectation = "an embedded ppt/media/* part with an image/* content type"
    elif kind == "audio":
        compatible = embedded_media and bool(
            content_type
            and (
                content_type.startswith("audio/")
                or content_type == "application/octet-stream"
            )
        )
        expectation = "an embedded audio part with a non-XML audio content type"
    elif kind == "video":
        compatible = embedded_media and bool(
            content_type
            and (
                content_type.startswith("video/")
                or content_type == "application/octet-stream"
            )
        )
        expectation = "an embedded video part with a non-XML video content type"
    elif kind == "media":
        compatible = embedded_media and bool(
            content_type
            and not content_type.endswith("+xml")
            and content_type not in {"application/xml", "text/xml"}
            and not content_type.startswith("text/")
        )
        expectation = "an embedded media part with a non-XML content type"
    else:
        return True, None
    return compatible, expectation


def _validate_relationships(
    package: zipfile.ZipFile,
    names: set[str],
    content_types: dict[str, str],
) -> tuple[dict[str, set[str]], list[str]]:
    references = {"slide": set(), "notes": set(), "media": set(), "image": set()}
    errors: list[str] = []
    for rels_path in sorted(name for name in names if name.endswith(".rels")):
        try:
            source, source_dir = _relationship_source(rels_path)
        except ValueError as error:
            errors.append(_normalize_controls(error))
            continue
        if source is not None and source not in names:
            errors.append(
                f"Relationships part '{rels_path}' has missing source part '{source}'."
            )
        try:
            root = ET.fromstring(package.read(rels_path))
        except (ET.ParseError, KeyError, OSError, zipfile.BadZipFile) as error:
            errors.append(
                f"Relationships part '{rels_path}' is unreadable: "
                f"{type(error).__name__}: {_normalize_controls(error)}"
            )
            continue
        if root.tag != f"{{{RELATIONSHIPS_NS}}}Relationships":
            errors.append(f"Relationships part '{rels_path}' has an invalid root element.")
            continue

        for relationship in root:
            if relationship.tag != f"{{{RELATIONSHIPS_NS}}}Relationship":
                continue
            rel_id = relationship.get("Id") or "[missing Id]"
            rel_type = relationship.get("Type") or ""
            target = relationship.get("Target") or ""
            mode = (relationship.get("TargetMode") or "Internal").lower()
            if not rel_type or not target:
                errors.append(
                    f"Relationship '{rel_id}' in '{rels_path}' lacks Type or Target."
                )
                continue
            kind = rel_type.rsplit("/", 1)[-1]
            if mode == "external":
                if kind != "hyperlink":
                    errors.append(
                        f"External asset relationship '{rel_id}' in '{rels_path}' "
                        f"targets '{_normalize_controls(target)}'; embed the asset."
                    )
                continue
            if mode != "internal":
                errors.append(
                    f"Relationship '{rel_id}' in '{rels_path}' has invalid "
                    f"TargetMode '{_normalize_controls(mode)}'."
                )
                continue
            try:
                resolved = _resolve_internal_target(source_dir, target)
            except ValueError as error:
                errors.append(_normalize_controls(error))
                continue
            if resolved not in names:
                errors.append(
                    f"Relationship '{rel_id}' in '{rels_path}' targets missing "
                    f"package part '{resolved}'."
                )
            compatible, expectation = _relationship_target_compatible(
                kind, resolved, content_types.get(resolved)
            )
            if not compatible:
                errors.append(
                    f"{kind} relationship '{rel_id}' in '{rels_path}' has "
                    f"incompatible target '{resolved}' with content type "
                    f"'{content_types.get(resolved) or '[missing]'}'; expected "
                    f"{expectation}."
                )
            if compatible and kind == "slide" and source == "ppt/presentation.xml":
                references["slide"].add(resolved)
            elif (
                compatible
                and kind == "notesSlide"
                and source is not None
                and SLIDE_PART.fullmatch(source)
            ):
                references["notes"].add(resolved)
            elif (
                kind in {"image", "audio", "video", "media"}
                and resolved in names
                and resolved.startswith("ppt/media/")
            ):
                references["media"].add(resolved)
                if kind == "image":
                    references["image"].add(resolved)
    return references, errors


def _validate_orphans(names: set[str], references: dict[str, set[str]]) -> list[str]:
    errors: list[str] = []
    categories = (
        ("slide", sorted(name for name in names if SLIDE_PART.fullmatch(name))),
        ("notes", sorted(name for name in names if NOTES_PART.fullmatch(name))),
        ("media", sorted(name for name in names if name.startswith("ppt/media/"))),
    )
    for kind, parts in categories:
        for part in parts:
            if part not in references[kind]:
                errors.append(f"Orphaned {kind} package part '{part}' is not referenced.")
    return errors


def _validate_images(
    package: zipfile.ZipFile,
    media_parts: Iterable[str],
    content_types: dict[str, str],
) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    image_parts = sorted(set(media_parts))
    raster_parts = [
        part
        for part in image_parts
        if PurePosixPath(part).suffix.lower() in RASTER_EXTENSIONS
    ]
    warnings.extend(
        f"Image package part '{part}' uses a format not decoded by Pillow; "
        "inspect it manually."
        for part in image_parts
        if part not in raster_parts
    )
    if not raster_parts:
        return errors, warnings
    try:
        from PIL import Image, UnidentifiedImageError
    except ImportError:
        warnings.append(
            "Pillow is unavailable; referenced raster image blobs were not decoded."
        )
        return errors, warnings

    for part in raster_parts:
        try:
            blob = package.read(part)
            with Image.open(BytesIO(blob)) as image:
                image.verify()
        except (KeyError, OSError, SyntaxError, UnidentifiedImageError) as error:
            errors.append(
                f"Image package part '{part}' is not a valid raster image: "
                f"{type(error).__name__}: {_normalize_controls(error)}"
            )
    return errors, warnings


def _inspect_package(path: Path, result: dict[str, Any]) -> bool:
    try:
        package = zipfile.ZipFile(path)
    except (OSError, zipfile.BadZipFile):
        result["errors"].append("Input is not a readable ZIP-based PPTX package.")
        return False

    with package:
        infos = [entry for entry in package.infolist() if not entry.is_dir()]
        names_list = [entry.filename for entry in infos]
        names = set(names_list)
        result["slides"] = sum(bool(SLIDE_PART.fullmatch(name)) for name in names)
        result["media"] = sum(name.startswith("ppt/media/") for name in names)
        result["notes_slides"] = sum(
            bool(NOTES_PART.fullmatch(name)) for name in names
        )

        duplicate_names = sorted(
            name for name, count in Counter(names_list).items() if count > 1
        )
        if duplicate_names:
            result["errors"].extend(
                f"Duplicate ZIP package name '{_normalize_controls(name)}' is not allowed."
                for name in duplicate_names
            )
        result["errors"].extend(_archive_limit_errors(infos))
        if result["errors"]:
            return False

        try:
            bad_crc = package.testzip()
        except (OSError, RuntimeError, zipfile.BadZipFile) as error:
            result["errors"].append(
                f"ZIP CRC validation failed: {type(error).__name__}: "
                f"{_normalize_controls(error)}"
            )
            return False
        if bad_crc is not None:
            result["errors"].append(
                f"ZIP CRC validation failed for package part "
                f"'{_normalize_controls(bad_crc)}'."
            )
            return False

        try:
            defaults, overrides = _content_types(package)
        except (ET.ParseError, KeyError, OSError, ValueError, zipfile.BadZipFile) as error:
            result["errors"].append(
                "The PPTX cannot be reopened because [Content_Types].xml is "
                "missing or unreadable: "
                f"{type(error).__name__}: {_normalize_controls(error)}"
            )
            return False
        content_types, content_errors = _validate_content_type_coverage(
            names, defaults, overrides
        )
        references, relationship_errors = _validate_relationships(
            package, names, content_types
        )
        result["errors"].extend(content_errors)
        result["errors"].extend(relationship_errors)
        result["errors"].extend(_validate_orphans(names, references))
        image_errors, image_warnings = _validate_images(
            package,
            references["image"],
            content_types,
        )
        result["errors"].extend(image_errors)
        result["warnings"].extend(image_warnings)
    return True


def _notes_message(slides: list[int]) -> str:
    label = "Slide" if len(slides) == 1 else "Slides"
    verb = "has" if len(slides) == 1 else "have"
    numbers = ", ".join(str(number) for number in slides)
    return f"{label} {numbers} {verb} no meaningful speaker notes."


def _compose(outer: Matrix, inner: Matrix) -> Matrix:
    oa, ob, oc, od, oe, of = outer
    ia, ib, ic, id_, ie, iff = inner
    return (
        oa * ia + oc * ib,
        ob * ia + od * ib,
        oa * ic + oc * id_,
        ob * ic + od * id_,
        oa * ie + oc * iff + oe,
        ob * ie + od * iff + of,
    )


def _apply(matrix: Matrix, x: float, y: float) -> tuple[float, float]:
    a, b, c, d, e, f = matrix
    return a * x + c * y + e, b * x + d * y + f


def _rotation(degrees: float, center_x: float, center_y: float) -> Matrix:
    radians = math.radians(degrees % 360)
    cosine = math.cos(radians)
    sine = math.sin(radians)
    return (
        cosine,
        sine,
        -sine,
        cosine,
        center_x - cosine * center_x + sine * center_y,
        center_y - sine * center_x - cosine * center_y,
    )


def _shape_geometry(shape: Any) -> tuple[float, float, float, float, float]:
    return (
        float(shape.left),
        float(shape.top),
        float(shape.width),
        float(shape.height),
        float(shape.rotation or 0),
    )


def _visual_aabb(shape: Any, parent: Matrix) -> tuple[float, float, float, float]:
    left, top, width, height, rotation = _shape_geometry(shape)
    local_rotation = _rotation(rotation, left + width / 2, top + height / 2)
    transform = _compose(parent, local_rotation)
    points = (
        _apply(transform, left, top),
        _apply(transform, left + width, top),
        _apply(transform, left + width, top + height),
        _apply(transform, left, top + height),
    )
    xs = [point[0] for point in points]
    ys = [point[1] for point in points]
    return min(xs), min(ys), max(xs), max(ys)


def _group_child_transform(group: Any, parent: Matrix) -> Matrix:
    left, top, width, height, rotation = _shape_geometry(group)
    child_x, child_y, child_width, child_height = _group_child_coordinates(group)
    if child_width <= 0 or child_height <= 0:
        raise ValueError("group child coordinate extent is zero or negative")
    scale_x = width / child_width
    scale_y = height / child_height
    child_map: Matrix = (
        scale_x,
        0.0,
        0.0,
        scale_y,
        left - child_x * scale_x,
        top - child_y * scale_y,
    )
    group_rotation = _rotation(rotation, left + width / 2, top + height / 2)
    return _compose(parent, _compose(group_rotation, child_map))


def _group_child_coordinates(group: Any) -> tuple[float, float, float, float]:
    """Isolate the python-pptx 1.0.2 private group transform interface."""
    child_offset = group._element.chOff
    child_extent = group._element.chExt
    return (
        float(child_offset.x),
        float(child_offset.y),
        float(child_extent.cx),
        float(child_extent.cy),
    )


def _overflow_finding(
    slide_number: int,
    shape: Any,
    parent: Matrix,
    slide_width: int,
    slide_height: int,
) -> dict[str, Any] | None:
    minimum_x, minimum_y, maximum_x, maximum_y = _visual_aabb(shape, parent)
    overflows = {
        "left": max(0.0, -minimum_x),
        "right": max(0.0, maximum_x - slide_width),
        "top": max(0.0, -minimum_y),
        "bottom": max(0.0, maximum_y - slide_height),
    }
    overflows = {side: amount for side, amount in overflows.items() if amount > 0}
    if not overflows:
        return None

    raw_percentages = {
        side: amount
        / (slide_width if side in {"left", "right"} else slide_height)
        * 100
        for side, amount in overflows.items()
    }
    percentages = {
        side: round(percent, 6) for side, percent in raw_percentages.items()
    }
    wholly_outside = (
        maximum_x <= 0
        or minimum_x >= slide_width
        or maximum_y <= 0
        or minimum_y >= slide_height
    )
    severity = (
        "error"
        if wholly_outside
        or any(percent > BLEED_LIMIT_PERCENT for percent in raw_percentages.values())
        else "warning"
    )
    action = (
        "Shape is wholly outside the slide canvas; move or resize it into view."
        if wholly_outside
        else "Move or resize the shape so it stays within the slide canvas."
    )
    return {
        "slide": slide_number,
        "shape": _normalize_controls(shape.name or f"shape {shape.shape_id}"),
        "shape_id": int(shape.shape_id),
        "severity": severity,
        "sides": list(percentages),
        "overflow_percent": percentages,
        "action": action,
    }


def _finding_message(finding: dict[str, Any]) -> str:
    details = ", ".join(
        f"{side} {finding['overflow_percent'][side]:.6g}%"
        for side in finding["sides"]
    )
    if finding["severity"] == "warning":
        return (
            f"Slide {finding['slide']} shape '{finding['shape']}' has possible "
            f"intentional bleed ({details}); verify it visually."
        )
    return (
        f"Slide {finding['slide']} shape '{finding['shape']}' is materially "
        f"outside the canvas ({details}). {finding['action']}"
    )


def _inspect_shape(
    slide_number: int,
    shape: Any,
    parent: Matrix,
    slide_width: int,
    slide_height: int,
    result: dict[str, Any],
) -> None:
    name = _normalize_controls(getattr(shape, "name", "unknown shape"))
    left, top, width, height, _ = _shape_geometry(shape)
    shape_type = getattr(getattr(shape, "shape_type", None), "name", "")
    is_connector = shape_type == "LINE"
    is_group = shape_type == "GROUP" and hasattr(shape, "shapes")
    zero_area = width <= 0 or height <= 0
    allowed_connector = is_connector and ((width == 0) ^ (height == 0)) and max(width, height) > 0
    if zero_area and not allowed_connector:
        result["errors"].append(
            f"Slide {slide_number} shape '{name}' has zero-area or negative geometry; "
            "resize or remove it."
        )
    else:
        finding = _overflow_finding(
            slide_number, shape, parent, slide_width, slide_height
        )
        if finding is not None:
            result["out_of_bounds"].append(finding)
            message = _finding_message(finding)
            if finding["severity"] == "error":
                result["errors"].append(message)
            else:
                result["warnings"].append(message)

    if is_group:
        try:
            child_transform = _group_child_transform(shape, parent)
        except (AttributeError, TypeError, ValueError, ZeroDivisionError) as error:
            result["errors"].append(
                f"Slide {slide_number} group '{name}' has invalid child geometry: "
                f"{_normalize_controls(error)}"
            )
            return
        for child in shape.shapes:
            _inspect_shape(
                slide_number,
                child,
                child_transform,
                slide_width,
                slide_height,
                result,
            )


def _inspect_presentation(
    path: Path,
    result: dict[str, Any],
    *,
    require_notes: bool,
) -> tuple[int, int] | None:
    try:
        from pptx import Presentation

        presentation = Presentation(path)
    except Exception as error:  # python-pptx exposes several package exceptions
        result["errors"].append(
            "The PPTX package could not be reopened as a presentation: "
            f"{type(error).__name__}: {_normalize_controls(error)}"
        )
        return None

    dimensions: tuple[int, int] | None = None
    try:
        raw_width = presentation.slide_width
        raw_height = presentation.slide_height
        if raw_width is None or raw_height is None:
            result["errors"].append("Presentation slide dimensions are missing.")
            return None
        slide_width = int(raw_width)
        slide_height = int(raw_height)
        if slide_width <= 0 or slide_height <= 0:
            result["errors"].append(
                "Presentation slide dimensions must both be positive."
            )
            return None
        dimensions = (slide_width, slide_height)
        if len(presentation.slides) != result["slides"]:
            result["errors"].append(
                "Package slide count does not match the reopened presentation "
                f"({result['slides']} package parts vs "
                f"{len(presentation.slides)} slides)."
            )

        missing_notes: list[int] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            meaningful_notes = False
            if slide.has_notes_slide:
                text_frame = slide.notes_slide.notes_text_frame
                meaningful_notes = bool(
                    text_frame is not None and text_frame.text.strip()
                )
            if not meaningful_notes:
                missing_notes.append(slide_number)
            for shape in slide.shapes:
                _inspect_shape(
                    slide_number,
                    shape,
                    IDENTITY,
                    slide_width,
                    slide_height,
                    result,
                )

        if missing_notes:
            message = _notes_message(missing_notes)
            if require_notes:
                result["errors"].append(message)
            else:
                result["warnings"].append(message)
    except Exception as error:
        result["errors"].append(
            "Presentation structural traversal failed: "
            f"{type(error).__name__}: {_normalize_controls(error)}"
        )
    return dimensions


def inspect_pptx(
    path: Path, *, require_notes: bool = False
) -> tuple[dict[str, Any], tuple[int, int] | None]:
    """Return a stable structural result and slide dimensions in EMU."""
    result = _empty_result(path)
    try:
        if not path.exists():
            result["errors"].append(f"Input PPTX does not exist: {result['file']}")
            return result, None
        if not path.is_file():
            result["errors"].append(f"Input PPTX is not a file: {result['file']}")
            return result, None

        safe_to_reopen = _inspect_package(path, result)
        dimensions = (
            _inspect_presentation(path, result, require_notes=require_notes)
            if safe_to_reopen
            else None
        )
    except Exception as error:
        result["errors"].append(
            "PPTX structural inspection failed: "
            f"{type(error).__name__}: {_normalize_controls(error)}"
        )
        dimensions = None
    result["ok"] = not result["errors"]
    return result, dimensions


def _format_dimensions(dimensions: tuple[int, int] | None) -> str:
    if dimensions is None:
        return "unavailable"
    width, height = dimensions
    return (
        f"{width / EMU_PER_INCH:.2f} x {height / EMU_PER_INCH:.2f} in "
        f"({width} x {height} EMU)"
    )


def markdown_report(
    result: dict[str, Any], dimensions: tuple[int, int] | None
) -> str:
    lines = [
        "# Paper2PPT structural QA report",
        "",
        f"- File: {_markdown_escape(result['file'])}",
        f"- Status: {'pass' if result['ok'] else 'fail'}",
        f"- Slides: {result['slides']}",
        f"- Media parts: {result['media']}",
        f"- Notes slide parts: {result['notes_slides']}",
        f"- Slide dimensions: {_format_dimensions(dimensions)}",
        "",
        "This report covers structural checks only; content QA and rendered visual "
        "QA remain separate gates.",
        "",
        "## Bounds findings",
        "",
    ]
    if result["out_of_bounds"]:
        for finding in result["out_of_bounds"]:
            details = ", ".join(
                f"{side} {finding['overflow_percent'][side]:.6g}%"
                for side in finding["sides"]
            )
            lines.append(
                f"- Slide {finding['slide']}, {_markdown_escape(finding['shape'])} "
                f"({finding['severity']}): outside on {details}. "
                f"{_markdown_escape(finding['action'])}"
            )
    else:
        lines.append("- None.")

    for title, key in (("Warnings", "warnings"), ("Errors", "errors")):
        lines.extend(["", f"## {title}", ""])
        if result[key]:
            lines.extend(f"- {_markdown_escape(message)}" for message in result[key])
        else:
            lines.append("- None.")

    lines.extend(
        [
            "",
            "## Next action",
            "",
            (
                "- Proceed to content QA and rendered visual QA."
                if result["ok"]
                else "- Correct every error, regenerate the PPTX, and rerun this inspector."
            ),
            "",
        ]
    )
    return "\n".join(lines)


def _human_summary(result: dict[str, Any]) -> str:
    status = "PASS" if result["ok"] else "FAIL"
    lines = [
        f"Paper2PPT structural QA: {status}",
        f"File: {_normalize_controls(result['file'])}",
        f"Slides: {result['slides']}; media: {result['media']}; "
        f"notes parts: {result['notes_slides']}",
    ]
    lines.extend(f"Warning: {_normalize_controls(message)}" for message in result["warnings"])
    lines.extend(f"Error: {_normalize_controls(message)}" for message in result["errors"])
    return "\n".join(lines)


def _paths_identical(input_path: Path, report_path: Path) -> bool:
    try:
        if input_path.resolve(strict=False) == report_path.resolve(strict=False):
            return True
    except (OSError, RuntimeError):
        pass
    try:
        return input_path.exists() and report_path.exists() and os.path.samefile(
            input_path, report_path
        )
    except (OSError, ValueError):
        return False


def _write_temp_file(descriptor: int, text: str) -> None:
    with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
        handle.write(text)
        handle.flush()
        os.fsync(handle.fileno())


def _write_report(
    report_path: Path,
    input_path: Path,
    result: dict[str, Any],
    dimensions: tuple[int, int] | None,
) -> None:
    if _paths_identical(input_path, report_path):
        raise OSError("report path refers to the same file as the input PPTX")
    if not report_path.parent.is_dir():
        raise OSError(
            f"report parent directory does not exist: {report_path.parent}"
        )

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{report_path.name}.",
        suffix=".tmp",
        dir=report_path.parent,
    )
    temporary = Path(temporary_name)
    try:
        _write_temp_file(descriptor, markdown_report(result, dimensions))
        descriptor = -1
        os.replace(temporary, report_path)
    except Exception:
        if descriptor >= 0:
            try:
                os.close(descriptor)
            except OSError:
                pass
        try:
            temporary.unlink()
        except FileNotFoundError:
            pass
        raise


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Inspect an editable PPTX package for structural QA issues."
    )
    parser.add_argument("deck", type=Path, help="PPTX file to inspect")
    parser.add_argument(
        "--require-notes",
        action="store_true",
        help="Fail when any slide lacks meaningful speaker-note text.",
    )
    parser.add_argument(
        "--json", action="store_true", help="Print the stable result as JSON."
    )
    parser.add_argument(
        "--report", type=Path, help="Write a deterministic Markdown QA report."
    )
    args = parser.parse_args(argv)

    result, dimensions = inspect_pptx(
        args.deck, require_notes=args.require_notes
    )
    if args.report is not None:
        try:
            _write_report(args.report, args.deck, result, dimensions)
        except Exception as error:
            result["errors"].append(
                f"Could not write QA report: {type(error).__name__}: "
                f"{_normalize_controls(error)}"
            )
            result["ok"] = False

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(_human_summary(result))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
